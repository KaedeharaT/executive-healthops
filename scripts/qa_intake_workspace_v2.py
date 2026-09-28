"""Chromium acceptance through public navigation, using an isolated synthetic member."""
import json
import sys
import re
from pathlib import Path
from playwright.sync_api import sync_playwright, expect

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'.runtime/intake-workspace-v2'
OUT=ROOT/'docs/images/intake-workspace-v2'
WORK.mkdir(parents=True,exist_ok=True)
OUT.mkdir(parents=True,exist_ok=True)

if '--prepare' in sys.argv:
    import sqlite3
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from executive_health_ai.models import Patient
    from executive_health_ai.services.management_workflow import ManagementWorkflowService
    with sqlite3.connect(ROOT/'executive_health_ai.db') as source,sqlite3.connect(WORK/'qa.db') as destination:source.backup(destination)
    with Session(create_engine('sqlite:///'+(WORK/'qa.db').as_posix())) as session:
        patient=Patient(display_name='Demo Intake Workspace V2',timezone='Asia/Tokyo');session.add(patient);session.flush()
        ManagementWorkflowService().start_intake(session,patient.id,2026,'验收健管');session.commit()
    (WORK/'manifest.json').write_text(json.dumps({'intake-v2':{'source':str(ROOT),'database':str(WORK/'qa.db'),'port':18583}}),encoding='utf-8')
    sys.exit(0)

def fixtures():
    from docx import Document
    questionnaire={'responses':{'家族健康史':[{'疾病类别':'心血管','具体疾病':'高血压','患病家属':'父亲','备注':'会员自述，待核对'}],
        '当前用药 / 营养补充':[{'名称':'维生素D','剂量':'1','单位':'片','频次':'每日','途径':'口服','开始日期':'2026-01-01','处方来源':'会员自述'}]}}
    (WORK/'questionnaire.json').write_text(json.dumps(questionnaire,ensure_ascii=False),encoding='utf-8')
    doc=Document();doc.add_paragraph('睡眠：七小时');doc.add_paragraph('本人关注：改善睡眠');doc.save(WORK/'history.docx')
    (WORK/'checkup-report.txt').write_text('体检日期：2026-09-28\nbirth_date: 1980-01-01\n体重 78 kg\nLDL-C 3.8 mmol/L\n',encoding='utf-8')


fixtures()


def persisted_results():
    from sqlalchemy import create_engine,select,func
    from sqlalchemy.orm import Session
    from executive_health_ai.models import Patient,AgentGoal,ReportExtractionCandidate,Observation,RiskEvent,HealthProblem,MedicationPlan
    with Session(create_engine('sqlite:///'+(WORK/'qa.db').as_posix())) as session:
        patient=session.scalar(select(Patient).where(Patient.display_name=='Demo Intake Workspace V2'))
        goals=list(session.scalars(select(AgentGoal).where(AgentGoal.member_id==patient.id)))
        rows=list(session.scalars(select(ReportExtractionCandidate).where(ReportExtractionCandidate.patient_id==patient.id)))
        actual={'goals':len(goals),'all_goals_completed':all(g.status=='COMPLETED' for g in goals),
            'candidate_count':len(rows),'measurement_candidates':sum(r.candidate_type=='OBSERVATION' for r in rows),
            'profile_candidates':sum(r.candidate_type=='PROFILE_FACT' for r in rows),
            'clinical_writes':{model.__name__:session.scalar(select(func.count(model.id)).where(model.patient_id==patient.id)) for model in (Observation,RiskEvent,HealthProblem,MedicationPlan)}}
        assert actual['goals']==3 and actual['all_goals_completed']
        assert actual['candidate_count']==16 and actual['measurement_candidates']==2
        assert not any(actual['clinical_writes'].values())
        return actual


with sync_playwright() as p:
    browser=p.chromium.launch(headless=True)
    page=browser.new_page(viewport={'width':1440,'height':1100})
    page.set_default_timeout(30000);errors=[];page.on('pageerror',lambda error:errors.append(str(error)))
    def settle():
        page.wait_for_timeout(600)
        page.locator('[data-testid="stStatusWidget"]').wait_for(state='hidden',timeout=120000)
        assert not page.locator('[data-testid="stException"]').count(),page.locator('body').inner_text()
    def button(name):page.get_by_role('button',name=name,exact=True).click();settle()
    def radio(name):page.get_by_role('radio',name=name,exact=True).locator('xpath=ancestor::label').click();settle()
    def shot(name):
        whole=name in {'01-health-record-page.png','08-no-duplicate-table.png'}
        if whole:
            page.set_viewport_size({'width':1440,'height':2700})
            page.locator('[data-testid="stMain"]').evaluate('(element)=>element.scrollTo(0,0)')
        page.screenshot(path=str(OUT/name),full_page=True)
        if whole:page.set_viewport_size({'width':1440,'height':1100})
    try:
        page.goto('http://127.0.0.1:18583',wait_until='networkidle');settle()
        if page.get_by_role('button',name='进入 HealthOps 运营后台',exact=True).count():button('进入 HealthOps 运营后台')
        radio('会员');search=page.get_by_label('搜索成员',exact=True);search.fill('Demo Intake Workspace V2');search.press('Enter');settle()
        page.locator('[data-testid="stDataFrame"]').last.click(position={'x':100,'y':55});settle()
        button('查看会员 / 进入Member360');radio('健康档案')
        expect(page.locator('input[type=file]')).to_have_count(1,timeout=30000)
        body=page.locator('body').inner_text()
        assert '资料用途' not in body and '辅助填写初始健康评估' not in body
        assert page.locator('.st-key-intake-agent-board').is_visible()
        assert page.locator('.st-key-intake-section-cards button').count()==10
        assert not page.locator('.st-key-soft-archive-details').count()
        shot('01-health-record-page.png')
        page.locator('.st-key-intake-upload-area').screenshot(path=str(OUT/'02-single-upload.png'))
        page.locator('input[type=file]').set_input_files([str(WORK/name) for name in ('checkup-report.txt','questionnaire.json','history.docx')]);settle()
        page.get_by_role('button',name='上传并整理资料',exact=True).click()
        expect(page.get_by_role('heading',name='健康管理助手正在整理资料',exact=True)).to_be_visible(timeout=60000)
        shot('03-agent-running.png')
        expect(page.get_by_role('heading',name='本次资料整理已完成',exact=True)).to_be_visible(timeout=120000);settle()
        shot('04-agent-completed.png')
        cards=page.locator('.st-key-intake-section-cards')
        expect(cards.get_by_role('button',name='家族史')).to_contain_text('已预填')
        expect(cards.get_by_role('button',name='用药')).to_contain_text('已预填')
        cards.screenshot(path=str(OUT/'05-clickable-sections.png'))
        cards.get_by_role('button',name='家族史').click();settle()
        expect(page.get_by_role('combobox',name='填写步骤')).to_have_value('2. 家族健康史')
        page.get_by_role('combobox',name='填写步骤').scroll_into_view_if_needed();shot('06-family-history-step.png')
        button('← 返回健康档案')
        medication=page.locator('.st-key-intake-section-cards').get_by_role('button',name='用药')
        medication.focus();page.keyboard.press('Enter');settle()
        expect(page.get_by_role('combobox',name='填写步骤')).to_have_value(re.compile('6. 当前用药'))
        page.get_by_role('combobox',name='填写步骤').scroll_into_view_if_needed();shot('07-medication-step.png')
        button('← 返回健康档案');shot('08-no-duplicate-table.png')
        assert not page.locator('.st-key-soft-archive-details').count()
        button('查看完整健康档案')
        expect(page.get_by_role('heading',name='完整健康档案',exact=True)).to_be_visible()
        assert page.locator('.st-key-soft-archive-details [data-testid="stDataFrame"]').count()==1
        shot('09-full-archive.png');button('← 返回健康档案主页')
        button('继续填写')
        expect(page.get_by_role('combobox',name='填写步骤')).to_have_value('1. 基础资料')
        # Confirm original report measurements without promoting medical facts.
        page.get_by_text('资料来源与待核对内容',exact=True).click();settle()
        page.get_by_label('人工核对与补充处理记录',exact=True).fill('已核对原始体检测量；保留正式档案更新候选，初评提交不写入医疗事实')
        page.get_by_label('人工核对与补充处理记录',exact=True).press('Tab');settle()
        button('记录人工处理')
        # Confirm sources through the real wizard; continuation must advance.
        page.get_by_role('checkbox',name='已核对本步骤原始资料，确认当前填写或保留未知').locator('xpath=ancestor::label').click()
        button('保存草稿并继续');button('← 返回健康档案');button('继续填写')
        expect(page.get_by_role('combobox',name='填写步骤')).to_have_value('2. 家族健康史')
        for _ in range(9):
            box=page.get_by_role('checkbox',name='已核对本步骤原始资料，确认当前填写或保留未知')
            if not box.is_checked():box.locator('xpath=ancestor::label').click()
            button('保存草稿并继续')
        expect(page.get_by_role('combobox',name='填写步骤')).to_have_value('11. 确认提交')
        page.get_by_role('checkbox',name='我确认已逐项核对；未知项由健康管理团队继续确认').locator('xpath=ancestor::label').click();settle()
        button('提交初始评估');expect(page.get_by_text('问卷已提交',exact=True)).to_be_visible()
        button('← 返回健康档案')
        page.locator('.st-key-intake-agent-board').screenshot(path=str(OUT/'10-submitted-board.png'))
        page.set_viewport_size({'width':1440,'height':2600})
        page.locator('[data-testid="stMain"]').evaluate('(element)=>element.scrollTo(0,0)')
        shot('11-complete-workspace.png')
        assert not errors
        (OUT/'browser-results.json').write_text(json.dumps({'browser':browser.version,'navigation':'会员 → Member360 → 健康档案','upload_count':1,'files':3,'purpose_selector':False,'board_visible':True,'running':True,'completed':True,'cards':10,'family_step':2,'medication_step':6,'keyboard_enter':True,'first_incomplete_before_save':1,'first_incomplete_after_save':2,'submitted':True,'completed_board_retained':True,'duplicate_table':False,'full_archive':True,'AI':'本次未使用','knowledge':'本次未使用','persisted_results':persisted_results(),'errors':errors},ensure_ascii=False,indent=2),encoding='utf-8')
        print('Intake workspace Chromium acceptance passed',flush=True)
    except Exception:
        shot('failure.png');(OUT/'failure.txt').write_text(page.locator('body').inner_text(),encoding='utf-8');raise
    finally:browser.close()
