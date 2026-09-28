"""Real Chromium acceptance, synthetic DB only; no mocked percentage or delay."""
import io,json,os,re,sys,time
from pathlib import Path
from playwright.sync_api import sync_playwright,expect

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/images/v7-redesign/after';OUT.mkdir(parents=True,exist_ok=True)
URL='http://127.0.0.1:18592'
if URL not in {'http://127.0.0.1:18592'}:raise SystemExit('Only isolated progress QA instances are allowed')
if URL.endswith('18590'):OUT=OUT/'timeout';OUT.mkdir(exist_ok=True)
if '--retry' in sys.argv:OUT=OUT/'retry';OUT.mkdir(exist_ok=True)


with sync_playwright() as p:
    browser=p.chromium.launch(headless=True)
    page=browser.new_page(viewport={'width':1440,'height':900})
    def settle():
        page.wait_for_timeout(700)  # Let Streamlit receive the click before checking its script state.
        expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state','notRunning',timeout=90000)
        page.wait_for_function('!document.querySelector(\'[data-stale="true"]\')')
        assert not page.locator('[data-testid="stException"]').count()
    def button(name):page.get_by_role('button',name=name,exact=True).click();settle()
    def radio(name):page.get_by_role('radio',name=name,exact=True).locator('xpath=ancestor::label').click();settle()
    def check(name):
        target=page.get_by_role('checkbox',name=name,exact=True)
        if not target.is_checked():target.locator('xpath=ancestor::label').click();settle()
    def shot(name):
        page.locator('[data-testid="stMain"]').evaluate('(e)=>e.scrollTo(0,0)')
        page.screenshot(path=str(OUT/(name+'.png')))
        (OUT/(name+'.txt')).write_text(page.locator('body').inner_text(),encoding='utf-8')
        print(name,flush=True)
    try:
        page.goto(URL,wait_until='networkidle');settle()
        if page.get_by_role('button',name='进入 HealthOps 运营后台',exact=True).count():button('进入 HealthOps 运营后台')
        radio('会员')
        page.get_by_label('搜索成员',exact=True).fill('V6入口验证会员');page.get_by_label('搜索成员',exact=True).press('Enter');settle()
        for x in (160,100,15):
            page.locator('[data-testid="stDataFrame"]').last.click(position={'x':x,'y':55});settle()
            try:
                page.get_by_role('radio',name='健康档案',exact=True).wait_for(timeout=3000);break
            except Exception:
                if x==15:raise
        radio('健康档案')
        if '--retry' in sys.argv:button('重新尝试')
        if '--upload' in sys.argv:
            from docx import Document
            d=Document();d.add_paragraph('报告日期：2026-09-28');d.add_paragraph('会员本人说：现在每周快走三次。')
            buf=io.BytesIO();d.save(buf)
            responses={'基础资料':{'display_name':'V6入口验证会员（合成）','birth_date':'1980-01-01','sex':'male'},
                '家族健康史':[{'患病家属':'父亲','具体疾病':'会员回答暂不清楚'}],
                '生活方式':{'睡眠':'七小时','烟草':'不吸烟','运动':'每周快走三次','饮酒':'不饮酒'},
                '会员重点关注':{'concern':'改善睡眠'}}
            files=[{'name':'01合成问卷.json','mimeType':'application/json','buffer':json.dumps({'responses':responses},ensure_ascii=False).encode()},
                   {'name':'02合成体检报告.txt','mimeType':'text/plain','buffer':'体检日期：2026-09-28\n体重 78 kg\n'.encode()},
                   {'name':'03合成既往档案.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':buf.getvalue()}]
            uploader=page.get_by_text('上传健康资料',exact=True)
            if not page.locator('input[type=file]').is_visible():uploader.first.click();settle()
            page.locator('input[type=file]').set_input_files(files);button('上传并整理资料')
        if '--finish' in sys.argv:
            start=page.get_by_role('button',name=re.compile('^处理剩余'))
            if start.count():start.click();settle();shot('06-intake-exceptions')
            for _ in range(60):
                if page.get_by_role('button',name='确认并完成初始健康评估',exact=True).count():break
                if page.get_by_role('textbox',name='填写会员实际回答',exact=True).count():
                    page.get_by_role('textbox',name='填写会员实际回答',exact=True).fill('会员本人明确回答暂不清楚（合成）');button('保存答案并处理下一项')
                elif page.get_by_role('button',name='确认并处理下一项',exact=True).count():button('确认并处理下一项')
                elif page.get_by_role('button',name='完成本份资料核对',exact=True).count():
                    check('已核对原文件；未确认的档案候选保留待后续处理');button('完成本份资料核对')
                else:raise AssertionError(page.locator('body').inner_text())
            optional=page.get_by_role('checkbox',name='这些可选资料尚未提供，继续保留未知',exact=True)
            if optional.count():check('这些可选资料尚未提供，继续保留未知')
            check('我已核对来源及整理结果，确认提交初评资料');button('确认并完成初始健康评估')
            expect(page.get_by_role('progressbar',name='整体业务进度')).to_have_attribute('aria-valuenow','100')
            assert not page.locator('.agent-action-spinner').count();shot('05-health-record-agent-completed')
        else:
            seen=[];ai_values=set();ai_captured=False;deadline=time.monotonic()+700
            while time.monotonic()<deadline:
                body=page.locator('body').inner_text()
                bar=page.get_by_role('progressbar',name='整体业务进度')
                if bar.count():
                    percent=bar.get_attribute('aria-valuenow')
                    if percent not in seen:seen.append(percent);shot('progress-'+percent)
                    if '本地AI ● 运行中' in body:
                        ai_values.add(percent)
                        if not ai_captured:shot('04-health-record-agent-running');ai_captured=True
                    if '文件信息提取：2 / 3' in body and not (OUT/'03-file-two-of-three.png').exists():shot('03-file-two-of-three')
                    if '整体进度 71%' in body and not page.locator('.agent-action-spinner').count():
                        shot('intake-waiting-manager');break
                    if 'AI服务响应超时' in body and not page.locator('.agent-action-spinner').count():
                        shot('05-timeout');break
                # Observation polling only; never changes application state/progress.
                page.wait_for_timeout(300)
            result={'browser':browser.version,'percentages_observed':seen,'during_real_ai':sorted(ai_values),
                    'spinner_stopped':not page.locator('.agent-action-spinner').count()}
            (OUT/'browser-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
            print(json.dumps(result,ensure_ascii=False),flush=True)
            assert result['spinner_stopped'],'Agent did not reach a human/terminal state in observation window'
        browser.close()
    except Exception:
        shot('failure');raise
