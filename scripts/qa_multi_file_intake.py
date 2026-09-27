"""Real Chromium, synthetic member, isolated SQLite copy; never alters live data."""
import json
import sqlite3
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'.runtime/multi-file-intake';OUT=ROOT/'docs/images/multi-file-intake'
WORK.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
if '--prepare' in sys.argv:
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from executive_health_ai.models import Patient
    from executive_health_ai.services.management_workflow import ManagementWorkflowService
    with sqlite3.connect(ROOT/'executive_health_ai.db') as source,sqlite3.connect(WORK/'qa.db') as destination:source.backup(destination)
    with Session(create_engine('sqlite:///'+(WORK/'qa.db').as_posix())) as session:
        p=Patient(display_name='Demo Multi File QA',timezone='Asia/Tokyo');session.add(p);session.flush()
        ManagementWorkflowService().start_intake(session,p.id,2026,'QA健管');session.commit()
    (WORK/'manifest.json').write_text(json.dumps({'intake':{'source':str(ROOT),'database':str(WORK/'qa.db'),'port':18560}}),encoding='utf-8')
    from docx import Document
    doc=Document();doc.add_paragraph('睡眠：七小时');doc.add_paragraph('本人关注：改善睡眠');doc.save(WORK/'history.docx')
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='生活方式';sheet.append(['字段','值']);sheet.append(['睡眠','八小时']);sheet.append(['运动','每周步行']);book.save(WORK/'questionnaire.xlsx')
    (WORK/'allergy.json').write_text(json.dumps({'responses':{'过敏史':[{'名称':'花粉','来源':'会员自述'}]}},ensure_ascii=False),encoding='utf-8')
    print('Prepared isolated synthetic intake');sys.exit(0)

from playwright.sync_api import sync_playwright,expect
with sync_playwright() as p:
    browser=p.chromium.launch(headless=True);page=browser.new_page(viewport={'width':1440,'height':900})
    page.set_default_timeout(30000);errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    def settle():
        page.wait_for_timeout(700)
        page.locator('[data-testid="stStatusWidget"]').wait_for(state='hidden',timeout=120000)
        assert not page.locator('[data-testid="stException"]').count(),page.locator('body').inner_text()
    def button(name):page.get_by_role('button',name=name,exact=True).click();settle()
    def radio(name):page.get_by_role('radio',name=name,exact=True).locator('xpath=ancestor::label').click();settle()
    def shot(name):page.screenshot(path=str(OUT/name),full_page=True)
    try:
        page.goto('http://127.0.0.1:18560',wait_until='networkidle');settle()
        if page.get_by_role('button',name='进入 HealthOps 运营后台',exact=True).count():button('进入 HealthOps 运营后台')
        radio('会员');field=page.get_by_label('搜索成员',exact=True);field.fill('Demo Multi File QA');field.press('Enter');settle()
        page.locator('[data-testid="stDataFrame"]').last.click(position={'x':100,'y':55});settle()
        button('查看会员 / 进入Member360');radio('健康档案')
        button('开始评估' if page.get_by_role('button',name='开始评估',exact=True).count() else '继续填写')
        shot('01-upload-entry.png')
        page.locator('input[type=file]').set_input_files([str(WORK/name) for name in ('history.docx','questionnaire.xlsx','allergy.json')]);settle()
        button('逐份解析并预填初评')
        for _ in range(50):
            page.wait_for_timeout(1000)
            if page.get_by_role('checkbox',name='已核对本步骤原始资料，确认当前填写或保留未知').is_enabled():break
        settle();shot('02-file-progress-and-status.png')
        for index in range(10):
            if index==6:
                expect(page.get_by_label('睡眠',exact=True)).to_have_value('')
                expect(page.get_by_label('运动',exact=True)).to_have_value('每周步行')
                page.get_by_label('睡眠',exact=True).fill('七小时')
                page.get_by_label('冲突处理 / 暂不采用说明（存在冲突时必填）',exact=True).fill('已联系本人核实，采用七小时；另一份资料暂不采用')
                page.get_by_text('本步骤资料核对 · 存在冲突',exact=True).scroll_into_view_if_needed()
                shot('03-conflict-and-source.png')
            if index==8:expect(page.get_by_label('会员自己最想改善什么',exact=True)).to_have_value('改善睡眠')
            box=page.get_by_role('checkbox',name='已核对本步骤原始资料，确认当前填写或保留未知')
            if not box.is_checked():box.locator('xpath=ancestor::label').click()
            button('保存草稿并继续')
        shot('04-review-before-submit.png')
        page.get_by_role('checkbox',name='我确认已逐项核对；未知项由健康管理团队继续确认').locator('xpath=ancestor::label').click();settle()
        button('提交初始评估')
        expect(page.get_by_text('问卷已提交',exact=True)).to_be_visible();shot('05-submitted.png')
        assert not errors
        (OUT/'browser-results.json').write_text(json.dumps({'browser':browser.version,'viewport':'1440x900','multi_file_upload':True,'source_review':True,'conflict_no_overwrite':True,'prefill':True,'submitted':True,'errors':errors},indent=2),encoding='utf-8')
        print('Multi-file Chromium acceptance passed')
    except Exception:
        shot('failure.png');(OUT/'failure.txt').write_text(page.locator('body').inner_text(),encoding='utf-8');raise
    finally:browser.close()
