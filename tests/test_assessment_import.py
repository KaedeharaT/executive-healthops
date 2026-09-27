"""Multi-file intake uses real persistence, source review and the existing Agent."""
import io
import json
from uuid import UUID
import pytest
from sqlalchemy import create_engine,select,func
from sqlalchemy.orm import Session
from executive_health_ai.models import Patient,AgentGoal,Observation,RiskEvent,MedicationPlan,ReportExtractionRun
from executive_health_ai.models.base import Base
from executive_health_ai.services.assessment_import import AssessmentImportService,DATA_STEPS
from executive_health_ai.services.management_workflow import ManagementWorkflowService,STEPS,TABLE_FIELDS
from executive_health_ai.services.profile_ingestion import ProfileIngestionService,candidates
from executive_health_ai.agent.supervisor import HealthOpsAgentSupervisor

service=AssessmentImportService();workflow=ManagementWorkflowService()

@pytest.fixture
def env(tmp_path,monkeypatch):
    engine=create_engine('sqlite://');Base.metadata.create_all(engine)
    monkeypatch.setattr(ProfileIngestionService,'storage_root',tmp_path)
    monkeypatch.setenv('LOCAL_LLM_ENABLED','false')
    with Session(engine,expire_on_commit=False) as session:
        patient=Patient(display_name='Synthetic Intake',timezone='Asia/Tokyo');session.add(patient);session.flush()
        row=workflow.start_intake(session,patient.id,2026,'QA')
        yield session,patient,row

def native(responses):return json.dumps({'responses':responses},ensure_ascii=False).encode()

def upload(env,files):
    session,patient,row=env
    results=service.upload_batch(session,patient.id,row.id,files,actor='QA')
    for result in results:
        if not result['goal_id']:continue
        goal=session.get(AgentGoal,UUID(result['goal_id']))
        for _ in range(6):
            if goal.status!='RUNNING':break
            HealthOpsAgentSupervisor().execute_next_step(session,goal.id)
    return service.project(session,patient.id,row.id),results

def test_native_duplicates_conflicts_and_no_clinical_writes(env):
    s,p,row=env
    view,_=upload(env,[('a.json',native({'生活方式':{'睡眠':'七小时','烟草':'无'}})),('b.json',native({'生活方式':{'睡眠':'七小时','烟草':'偶尔'}}))])
    assert all(g.status=='WAITING_MANAGER' for g in view['goals'])
    assert len(view['steps'])==10
    sleep=next(g for g in view['groups'] if g['field']=='睡眠')
    assert len(sleep['rows'])==2 and sleep['duplicate_count']==1
    assert view['prefill']['生活方式']=={'睡眠':'七小时'}
    assert next(x for x in view['steps'] if x['step']=='生活方式')['status']=='存在冲突'
    assert row.responses=={}
    from datetime import datetime,timezone
    from executive_health_ai.services.product_projection import ProductProjectionService
    work=ProductProjectionService().manager(s,datetime.now(timezone.utc))
    assert all('2 项资料' in item.reason for item in work.items if item.source_type=='profile_intake')
    assert all(not f['run'].llm_used for f in view['files'])
    for model in (Observation,RiskEvent,MedicationPlan):assert s.scalar(select(func.count(model.id)))==0

@pytest.mark.parametrize('kind',['docx','xlsx','csv','txt','pdf'])
def test_supported_formats_keep_source(env,kind):
    buf=io.BytesIO()
    if kind=='docx':
        from docx import Document
        doc=Document();doc.add_paragraph('睡眠：七小时');doc.save(buf)
    elif kind=='xlsx':
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='生活方式';sheet.append(['字段','值']);sheet.append(['睡眠','七小时']);book.save(buf)
    elif kind=='pdf':
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        writer=PdfWriter();page=writer.add_blank_page(width=600,height=800)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 700 Td (birth_date: 1980-01-01) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream);writer.write(buf)
    else:buf.write(('字段,值\n睡眠,七小时\n' if kind=='csv' else '睡眠：七小时').encode())
    view,_=upload(env,[('health.'+kind,buf.getvalue())])
    assert view['goals'][0].status=='WAITING_MANAGER',view['goals'][0].next_action
    assert view['groups']
    data=view['groups'][0]['rows'][0].structured_data_json
    assert data['source_filename']=='health.'+kind and data['source_locator']
    if kind=='xlsx':assert 'A2:B2' in data['source_locator']
    if kind=='pdf':assert '第 1 页' in data['source_locator']

def test_failed_file_does_not_block_next_and_duplicate_upload(env):
    view,results=upload(env,[('bad.docx',b'bad'),('good.json',native({'生活方式':{'睡眠':'七小时'}}))])
    assert results[0]['error'] and view['prefill']['生活方式']['睡眠']=='七小时'
    again,_=upload(env,[('good.json',native({'生活方式':{'睡眠':'七小时'}}))])
    assert len(again['goals'])==1

def test_unavailable_and_mixed_text_are_visible(env):
    view,_=upload(env,[('mixed.txt','睡眠：七小时\n我希望改善每天午后的疲劳感。'.encode())])
    assert view['prefill']['生活方式']['睡眠']=='七小时'
    assert view['files'][0]['needs_check']
    assert view['files'][0]['run'].llm_status=='UNAVAILABLE'

def test_real_semantic_attempt_has_bound_evidence(env,monkeypatch):
    class Client:
        def generate_structured(self,**kwargs):
            assert kwargs['task']=='parse_health_intake'
            return {'facts':[{'section':'会员重点关注','field':'concern','value':'改善午后疲劳','evidence':'希望改善午后疲劳'}]}
    monkeypatch.setattr('executive_health_ai.services.intake_extraction.LocalLLMClient',Client)
    view,_=upload(env,[('free.txt','睡眠：七小时\n希望改善午后疲劳'.encode())])
    assert view['prefill']['会员重点关注']['concern']=='改善午后疲劳'
    assert view['files'][0]['run'].llm_used

def test_review_and_submission_preserve_sources(env):
    s,p,row=env
    view,_=upload(env,[('q.json',native({'生活方式':{'睡眠':'七小时'}}))])
    for step in STEPS[:-1]:
        data=[] if step in TABLE_FIELDS else {'display_name':p.display_name} if step=='基础资料' else {'concern':'改善睡眠'} if step=='会员重点关注' else {'睡眠':'七小时'} if step=='生活方式' else {}
        workflow.save_intake(s,p.id,2026,step,data,'QA')
    with pytest.raises(ValueError,match='来源核对'):workflow.submit_intake(s,p.id,row.id,'QA')
    for step in DATA_STEPS:service.record_review(s,row,step,'QA')
    workflow.submit_intake(s,p.id,row.id,'QA')
    assert row.status=='SUBMITTED' and len(row.review['document_intake']['reviews'])==10
    assert view['goals'][0].status=='COMPLETED'
    assert all(c.status=='INTAKE_REVIEWED' for c in candidates(s,view['goals'][0]))
    from executive_health_ai.ui.pages.manager.assistant import completed_values
    assert '初评资料已核对并提交' in completed_values(view['goals'][0],p.display_name)
    assert view['goals'][0].context_json['intake_candidate_count']==1

def test_conflict_requires_note_and_new_file_invalidates_review(env):
    s,p,row=env
    view,_=upload(env,[('a.json',native({'生活方式':{'睡眠':'七小时'}})),('b.json',native({'生活方式':{'睡眠':'八小时'}}))])
    workflow.save_intake(s,p.id,2026,'生活方式',{'睡眠':'七小时'},'QA')
    with pytest.raises(ValueError,match='冲突处理'):service.record_review(s,row,'生活方式','QA')
    service.record_review(s,row,'生活方式','QA','已联系本人核实，采用七小时')
    assert next(x for x in service.project(s,p.id,row.id)['steps'] if x['step']=='生活方式')['status']=='已填写'
    view,_=upload(env,[('c.json',native({'生活方式':{'睡眠':'六小时'}}))])
    assert next(x for x in view['steps'] if x['step']=='生活方式')['status']=='存在冲突'

def test_draft_birth_date_and_exact_year_no_overwrite(env):
    s,p,row=env
    other=workflow.start_intake(s,p.id,2027,'QA')
    view,_=upload(env,[('p.json',native({'基础资料':{'birth_date':'1980-01-01'},'生活方式':{'睡眠':'七小时'}}))])
    assert p.birth_date is None and other.responses=={} and row.responses=={}
    assert view['prefill']['基础资料']['birth_date']=='1980-01-01'
    assert service.project(s,p.id,other.id)['goals']==[]

def test_dose_conflicts_bind_to_named_record(env):
    view,_=upload(env,[(name,native({'当前用药 / 营养补充':[{'名称':'自述药物','剂量':dose,'单位':'mg'}]})) for name,dose in [('a.json','5'),('b.json','10')]])
    groups=[g for g in view['groups'] if g['field']=='剂量']
    assert len(groups)==1 and groups[0]['conflict']
    assert '剂量' not in view['prefill']['当前用药 / 营养补充'][0]

def test_unknown_symptom_score_remains_unknown(env):
    s,p,row=env
    view,_=upload(env,[('symptoms.json',native({'专项症状评估':[{'症状':'疲劳','原始回答':'偶尔疲劳'}]}))])
    data=service.form_data(view,'专项症状评估')
    workflow.save_intake(s,p.id,2026,'专项症状评估',data,'QA')
    assert not row.responses['专项症状评估'][0].get('原始分数')

def test_unreadable_file_manual_handoff_not_fake_completed(env):
    s,p,row=env
    view,_=upload(env,[('empty.txt',b' ')])
    assert view['goals'][0].status=='ESCALATED'
    with pytest.raises(ValueError,match='未识别'):service.validate_submit(s,row)
    service.acknowledge_file(s,p.id,row.id,view['goals'][0].id,'QA','已人工核对空文件，无可填写信息')
    service.finish(s,row,'QA')
    assert view['goals'][0].status=='CANCELLED' and not view['goals'][0].success_criteria['parsed']

def test_cannot_bypass_review_to_write_formal_archive(env):
    s,p,row=env;view,_=upload(env,[('q.json',native({'个人病史':[{'疾病或问题':'会员自述历史情况'}]}))])
    with pytest.raises(ValueError,match='不能直接写入'):
        ProfileIngestionService().approve(s,view['goals'][0],{},actor='QA',role='HEALTH_MANAGER')

def test_archived_member_cannot_upload(env):
    from executive_health_ai.services.member_archive import MemberArchiveService
    s,p,row=env
    MemberArchiveService().archive(s,p.id,expected_name=p.display_name,confirmation_name=p.display_name,actor='QA',role='HEALTH_MANAGER',confirmed=True)
    results=service.upload_batch(s,p.id,row.id,[('q.json',native({'生活方式':{'睡眠':'七小时'}}))],actor='QA')
    assert results[0]['error'] and not results[0]['goal_id']

def test_ai_cannot_drop_negation_or_invent_source(env,monkeypatch):
    class Client:
        def generate_structured(self,**kwargs):
            return {'facts':[{'section':'个人病史','field':'疾病或问题','value':'糖尿病','evidence':'无糖尿病'},
                {'section':'生活方式','field':'睡眠','value':'八小时','evidence':'睡眠八小时'}]}
    monkeypatch.setattr('executive_health_ai.services.intake_extraction.LocalLLMClient',Client)
    view,_=upload(env,[('p.txt','运动：步行\n无糖尿病'.encode())])
    assert '个人病史' not in view['prefill'] and '睡眠' not in view['prefill']['生活方式']
    assert view['files'][0]['needs_check']
